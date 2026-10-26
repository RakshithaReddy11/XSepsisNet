"""Current-session report downloads for the doctor dashboard."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from docx import Document
from docx.shared import Pt
from fastapi import APIRouter
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, ConfigDict, Field

from backend.services.preprocessing_service import FEATURE_COLUMNS

router = APIRouter()


class ReportRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")

    patient_id: str = Field(min_length=1)
    patient_name: str | None = None
    admission: dict[str, Any] = Field(default_factory=dict)
    hours: list[dict[str, Any]] = Field(min_length=6, max_length=6)
    prediction: dict[str, Any]
    explanation: dict[str, Any]


@router.post("/reports/docx", response_class=StreamingResponse)
def create_docx_report(payload: ReportRequest) -> StreamingResponse:
    document = Document()
    document.add_heading("XSepsisNet Patient XAI Report", level=0)
    document.add_paragraph(
        "Research output only. This is not a clinical diagnosis or treatment recommendation."
    )
    document.add_heading("Patient and session", level=1)
    document.add_paragraph(f"Patient ID: {payload.patient_id}")
    if payload.patient_name:
        document.add_paragraph(f"Patient name: {payload.patient_name}")
    for key, value in payload.admission.items():
        if value not in (None, ""):
            document.add_paragraph(f"{key.replace('_', ' ').title()}: {value}")

    document.add_heading("Prediction", level=1)
    for key in (
        "raw_probability",
        "calibrated_probability",
        "classification_probability",
        "classification_threshold",
        "predicted_class",
        "risk_level",
    ):
        if key in payload.prediction:
            document.add_paragraph(f"{key.replace('_', ' ').title()}: {payload.prediction[key]}")

    document.add_heading("Top clinical SHAP contributions", level=1)
    table = document.add_table(rows=1, cols=4)
    for cell, text in zip(table.rows[0].cells, ("Feature", "Total SHAP", "Absolute SHAP", "Hour")):
        cell.text = text
    for item in payload.explanation.get("clinical_contributions", [])[:10]:
        cells = table.add_row().cells
        cells[0].text = str(item.get("feature", ""))
        cells[1].text = f"{float(item.get('total_shap', 0.0)):.6f}"
        cells[2].text = f"{float(item.get('absolute_shap', 0.0)):.6f}"
        cells[3].text = str(item.get("most_influential_hour", ""))

    document.add_heading("Six-hour clinical observations", level=1)
    observations = document.add_table(rows=1, cols=2)
    observations.rows[0].cells[0].text = "Hour"
    observations.rows[0].cells[1].text = "Recorded features"
    for index, hour in enumerate(payload.hours, start=1):
        row = observations.add_row().cells
        row[0].text = str(index)
        row[1].text = ", ".join(
            f"{feature}={hour[feature]}" for feature in FEATURE_COLUMNS if hour.get(feature) not in (None, "")
        )

    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(10)

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": "attachment; filename=xsepsisnet_patient_xai_report.docx"},
    )


__all__ = ["router"]
